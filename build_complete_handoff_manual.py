from pathlib import Path
import runpy

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
BASE = ROOT / "output" / "documents" / "Arduino_Uniform_Stock_Kiosk_Formal_Beginner_Guide.docx"
OUT = ROOT / "CUTTE_Kiosk_Complete_Build_and_Transfer_Manual.docx"
BRIEF = ROOT / "CUTTE_Kiosk_Master_Builder_Brief.md"

NAVY = "17365D"
PALE = "F3F6F9"
GRAY = "D9D9D9"
WHITE = "FFFFFF"
TEXT = RGBColor(31, 31, 31)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table):
    tbl_pr = table._tbl.tblPr
    group = tbl_pr.find(qn("w:tblBorders"))
    if group is None:
        group = OxmlElement("w:tblBorders")
        tbl_pr.append(group)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), GRAY)
        group.append(element)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_para(doc, text="", bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        first.bold = True
        first.font.name = "Arial"
        first.font.color.rgb = TEXT
        rest = p.add_run(text[len(bold_lead):])
        rest.font.name = "Arial"
        rest.font.color.rgb = TEXT
    else:
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.color.rgb = TEXT
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        run.font.name = "Arial"
        run.font.color.rgb = TEXT


def add_numbered(doc, items):
    for number, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"{number}. {item}")
        run.font.name = "Arial"
        run.font.color.rgb = TEXT


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    repeat_header(table.rows[0])
    for i, label in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, NAVY)
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            if row_index % 2:
                shade(cells[i], PALE)
            margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.color.rgb = TEXT
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_together = True
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = TEXT


def markdown_table(headers, rows):
    out = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    out.extend("| " + " | ".join(str(v) for v in row) + " |" for row in rows)
    return out


def build_markdown():
    uniform_rows = [
        (0, "MALE_POLO", "Male Polo", "0 to 5"),
        (1, "FEMALE_BLOUSE", "Female Blouse", "6 to 11"),
        (2, "MALE_PANTS", "Male Pants", "12 to 17"),
        (3, "FEMALE_SKIRT", "Female Skirt", "18 to 23"),
        (4, "T_SHIRT", "School T-Shirt (Unisex)", "24 to 29"),
        (5, "PE_UNIFORM", "PE Uniform (Unisex)", "30 to 35"),
    ]
    lines = [
        "# CUTTE Kiosk Master Builder Brief",
        "",
        "## Purpose",
        "",
        "Use this brief as the single source of truth when transferring the capstone to another account, AI builder, programmer, or group member. Build the described Grade 12 STEM prototype without silently expanding its scope.",
        "",
        "## Fixed System Decisions",
        "",
        "- The project is a school uniform stock availability kiosk.",
        "- The prototype does not use MySQL, Firebase, SQL Server, a cloud database, or another database server.",
        "- The Arduino Uno owns the stock logic and working stock array.",
        "- EEPROM preserves 36 quantities after power loss.",
        "- The ESP32 hosts or serves the kiosk webpage and bridges browser requests to the Arduino through serial communication.",
        "- The user interface uses HTML, CSS, and JavaScript.",
        "- Student checks are read only and never reduce stock.",
        "- Only a validated administrator update can change a quantity.",
        "- Formal Uniform requires Male or Female; School T-Shirt and PE Uniform are unisex and skip gender.",
        "- Available means 4 to 30, Low Stock means 1 to 3, and Out of Stock means 0.",
        "- Green, yellow, and red LEDs represent the three statuses, but status words and quantities must also be displayed.",
        "- A physical assistance button produces one local buzzer event and one assistance message per press.",
        "- The prototype does not include purchasing, payment, reservations, barcode scanning, RFID, official inventory integration, or automatic stock deduction.",
        "",
        "## System Architecture",
        "",
        "Student or administrator -> HTML CSS JavaScript kiosk -> ESP32 web server -> serial link -> Arduino Uno -> stock array and EEPROM -> LEDs buzzer and assistance button.",
        "",
        "The stock result returns through the reverse path from the Arduino to the ESP32 and then to the browser.",
        "",
        "## Component Responsibilities",
        "",
    ]
    lines += markdown_table(
        ["Component", "Responsibility"],
        [
            ("Kiosk computer", "Displays the browser interface and receives user input."),
            ("HTML", "Defines screens, controls, labels, and content."),
            ("CSS", "Provides readable and touchscreen friendly presentation."),
            ("JavaScript", "Manages selections, navigation, requests, validation, results, loading, and errors."),
            ("ESP32", "Runs the local web server and translates HTTP requests into serial commands."),
            ("Arduino Uno", "Owns quantities, stock rules, EEPROM storage, LED control, button reading, and buzzer control."),
            ("EEPROM", "Retains the last confirmed quantities after restart."),
            ("Administrator", "Verifies physical stock and submits accurate replacement quantities."),
        ],
    )
    lines += [
        "",
        "## Inventory Model",
        "",
        "There are six stock item rows and six sizes, producing exactly 36 stored quantities.",
        "",
        "Selection rule: Formal Uniform requires Male or Female. Male shows Male Polo and Male Pants. Female shows Female Blouse and Female Skirt. School T-Shirt and PE Uniform are unisex and must skip the gender question.",
        "",
        "Sizes must remain in this order: XS, S, M, L, XL, XXL. Their offsets are 0, 1, 2, 3, 4, and 5.",
        "",
    ]
    lines += markdown_table(["Item index", "Code", "Display name", "EEPROM address block"], uniform_rows)
    lines += [
        "",
        "Use the formula EEPROM address = item index times 6 plus size index.",
        "",
        "Example: School T-Shirt Medium uses item index 4 and size index 2, so its address is 4 times 6 plus 2, which equals 26.",
        "",
        "## Stock Rules",
        "",
        "- Quantity 4 to 30: AVAILABLE and green LED.",
        "- Quantity 1 to 3: LOW STOCK and yellow LED.",
        "- Quantity 0: OUT OF STOCK and red LED.",
        "- Negative, decimal, blank, unknown, or out of range values are invalid.",
        "- Turn all status LEDs off before turning on the one matching the current result.",
        "",
        "## Serial Communication Contract",
        "",
        "```text",
        "GET|T_SHIRT|M",
        "STOCK|12|AVAILABLE",
        "SET|T_SHIRT|M|15",
        "UPDATED|T_SHIRT|M|15",
        "ASSIST|REQUESTED",
        "ERROR|BAD_REQUEST",
        "ERROR|UNKNOWN_ITEM",
        "ERROR|UNKNOWN_SIZE",
        "ERROR|BAD_QUANTITY",
        "ERROR|TIMEOUT",
        "```",
        "",
        "Every command ends with a newline. Use the vertical bar as the field separator. The ESP32 must not report an update as successful until the Arduino returns UPDATED.",
        "",
        "## Recommended Web Routes",
        "",
    ]
    lines += markdown_table(
        ["Route", "Purpose", "Expected behavior"],
        [
            ("GET /", "Open kiosk", "Return the HTML CSS and JavaScript interface."),
            ("GET /check?item=CODE&size=SIZE", "Student stock check", "Send GET to Arduino and return a verified result."),
            ("POST /update", "Administrator update", "Require a valid prototype session or PIN, validate input, send SET, and wait for UPDATED."),
            ("GET /status", "Assistance or device status", "Return the latest local assistance state or controller status if implemented."),
        ],
    )
    lines += [
        "",
        "## Student Workflow",
        "",
        "1. Open Home and choose Check Availability.",
        "2. Select Formal Uniform, School T-Shirt, or PE Uniform.",
        "3. If Formal Uniform is selected, choose Male or Female, then choose the matching garment.",
        "4. If School T-Shirt or PE Uniform is selected, skip the gender question because the item is unisex.",
        "5. Select one size and review the confirmation.",
        "6. Show a loading state while JavaScript requests the ESP32.",
        "7. ESP32 validates the request and sends GET to the Arduino.",
        "8. Arduino returns quantity and status and activates the matching LED.",
        "9. Display the item, size, quantity, status words, and guidance.",
        "10. Check Another returns to category selection. Home clears all previous choices.",
        "11. A timeout or malformed response shows an error with Retry and Home and must never show a false stock status.",
        "",
        "## Administrator Workflow",
        "",
        "1. Open the separate administrator login screen.",
        "2. Enter the prototype PIN.",
        "3. Select a category and size. Use Male or Female only for Formal Uniform; skip gender for School T-Shirt and PE Uniform.",
        "4. Request and display the current Arduino quantity.",
        "5. Compare it with verified physical stock.",
        "6. Enter a replacement whole number within the allowed range.",
        "7. Validate in JavaScript and again on the Arduino.",
        "8. ESP32 sends SET only after validation succeeds.",
        "9. Arduino changes the correct array value and writes only the matching EEPROM address.",
        "10. Display success only after UPDATED is received.",
        "11. Check the item again and restart the Arduino during testing to verify persistence.",
        "",
        "## Assistance Workflow",
        "",
        "1. The button is normally HIGH with INPUT_PULLUP.",
        "2. A press changes the input from HIGH to LOW.",
        "3. Debounce the signal so one press creates one event.",
        "4. Sound the buzzer for a short controlled period.",
        "5. Send ASSIST|REQUESTED to the ESP32.",
        "6. Show Assistance Requested on the kiosk.",
        "7. Treat this as a local alert only. It does not send SMS, email, or internet notifications.",
        "",
        "## Hardware Wiring Contract",
        "",
    ]
    lines += markdown_table(
        ["Function", "Suggested Arduino pin", "Requirement"],
        [
            ("Assistance button", "D2", "Connect button to ground and use INPUT_PULLUP."),
            ("Green LED", "D6", "Use a suitable resistor such as 220 ohms."),
            ("Yellow LED", "D7", "Use a suitable resistor such as 220 ohms."),
            ("Red LED", "D8", "Use a suitable resistor such as 220 ohms."),
            ("Buzzer control", "D9", "Use a transistor driver if required by buzzer current."),
            ("Arduino serial RX", "D10", "Receive ESP32 TX at 3.3 volt logic."),
            ("Arduino serial TX", "D11", "Protect the ESP32 RX using verified level conversion."),
        ],
    )
    lines += [
        "",
        "The Arduino Uno uses 5 volt logic and the ESP32 uses 3.3 volt logic. Never connect the Uno TX output directly to an ESP32 input. Use a logic level converter or a verified voltage divider, and connect both boards to a common ground. Disconnect power before changing wiring.",
        "",
        "## Required Arduino Functions",
        "",
        "- initializeInventory for safe first run EEPROM setup.",
        "- loadInventory for restoring all 36 saved values.",
        "- findItemCode and findSizeCode for mapping text codes to indexes.",
        "- statusFor for applying the exact thresholds.",
        "- setStatusLeds for exclusive LED control.",
        "- processGet for read only checks.",
        "- processSet for validated updates and EEPROM writing.",
        "- processCommand for parsing complete newline terminated commands.",
        "- monitorAssistanceButton for debounce, buzzer timing, and the assistance event.",
        "",
        "## Required ESP32 Functions",
        "",
        "- Create a local access point or use a controlled school network only if approved.",
        "- Serve the kiosk webpage.",
        "- Validate required HTTP parameters.",
        "- Send newline terminated serial commands.",
        "- Wait for an Arduino reply with a finite timeout.",
        "- Return errors instead of inventing stock results.",
        "- Protect the administrator update route with at least the prototype PIN or session.",
        "- Show success only after the Arduino confirms the saved value.",
        "",
        "## Required Interface Screens",
        "",
        "- Home.",
        "- Category selection: Formal Uniform, School T-Shirt, or PE Uniform.",
        "- Gender selection shown only for Formal Uniform.",
        "- Formal garment selection filtered by Male or Female.",
        "- Size selection.",
        "- Confirmation.",
        "- Result.",
        "- Error and retry.",
        "- Administrator login.",
        "- Inventory management.",
        "",
        "Use large controls, readable text, plain labels, visible focus indicators, and status words in addition to color.",
        "",
        "## Build Order",
        "",
        "1. Test the three LEDs, button, and buzzer on the Arduino.",
        "2. Create the 6 by 6 stock array and print sample quantities in Serial Monitor.",
        "3. Implement and test GET commands.",
        "4. Implement SET validation in RAM before adding EEPROM.",
        "5. Add safe EEPROM initialization, loading, and single address updates.",
        "6. Establish ESP32 to Arduino serial communication with fixed test commands.",
        "7. Create the ESP32 web server and test the check route from a browser.",
        "8. Build the student interface using sample data, then connect it to the real route.",
        "9. Add the administrator interface and protected update route.",
        "10. Add assistance event monitoring and the kiosk message.",
        "11. Run the complete functional test plan and record actual results.",
        "12. Build the physical enclosure only after electronics and software work reliably on the bench.",
        "",
        "## Required Acceptance Tests",
        "",
        "- All six stock item rows and six sizes can be selected.",
        "- Formal Uniform requires Male or Female and shows only the matching garments.",
        "- School T-Shirt and PE Uniform never ask for gender.",
        "- All 36 combinations return their expected quantities.",
        "- Quantity 4 to 30 shows Available and only the green LED.",
        "- Quantity 1 to 3 shows Low Stock and only the yellow LED.",
        "- Quantity 0 shows Out of Stock and only the red LED.",
        "- Repeated student checks never change stock.",
        "- Valid administrator updates are confirmed and remain after restart.",
        "- Blank, negative, decimal, oversized, and unknown values are rejected without changing the old quantity.",
        "- A disconnected Arduino produces a timeout or controller error, not a stock result.",
        "- One button press produces one buzzer event and one assistance message.",
        "- Arduino to ESP32 voltage is safely converted and both devices share ground.",
        "- The test sheet contains dates, versions, actual results, pass or fail status, and corrective actions.",
        "",
        "## Features Outside the Current Scope",
        "",
        "Do not claim the following as completed: purchases, payments, online ordering, reservations, automatic sale deduction, barcode or RFID tracking, centralized database, cloud synchronization, official school inventory integration, multiple kiosk synchronization, staff SMS or email alerts, or production grade authentication.",
        "",
        "## Copy Paste Master Prompt for Another Builder",
        "",
        "```text",
        "Build a Grade 12 STEM capstone prototype named Development of an Arduino Based Interactive Kiosk for School Uniform Stock Availability. Preserve the following specification exactly unless the student team explicitly approves a change.",
        "",
        "Use an Arduino Uno as the stock and hardware controller, an ESP32 as the local web server and serial bridge, and an HTML CSS JavaScript interface on a kiosk computer or touchscreen. Do not add MySQL, Firebase, SQL Server, cloud storage, or another database. Store exactly 36 stock quantities as a 6 by 6 Arduino array and persist confirmed values in Arduino EEPROM. The item order is MALE_POLO, FEMALE_BLOUSE, MALE_PANTS, FEMALE_SKIRT, T_SHIRT, PE_UNIFORM. The size order is XS, S, M, L, XL, XXL. EEPROM address equals item index times 6 plus size index.",
        "",
        "The student first chooses Formal Uniform, School T-Shirt, or PE Uniform. Formal Uniform must ask for Male or Female. Male shows only Male Polo and Male Pants; Female shows only Female Blouse and Female Skirt. School T-Shirt and PE Uniform are unisex, so they must never ask for gender. Gender is an interface selection rule, not another EEPROM field or database dimension.",
        "",
        "A student stock check is read only and must never reduce quantity. Each item-and-size value has a hard maximum of 30. The status rules are: 4 to 30 is AVAILABLE with the green LED, 1 to 3 is LOW STOCK with the yellow LED, and 0 is OUT OF STOCK with the red LED. Display status words and quantity, not color alone. Turn all status LEDs off before activating the matching one.",
        "",
        "Use newline terminated serial commands with vertical bar separators. Support GET|ITEM|SIZE, STOCK|QUANTITY|STATUS, SET|ITEM|SIZE|QUANTITY, UPDATED|ITEM|SIZE|QUANTITY, ASSIST|REQUESTED, and clear ERROR responses. Validate every administrator update in both JavaScript and Arduino. Reject blank, negative, decimal, oversized, unknown item, and unknown size input. Keep the old value after any invalid update. Show success only after the Arduino replies UPDATED.",
        "",
        "Provide Home, category selection, conditional Formal gender selection, Formal garment selection, size selection, confirmation, result, error and retry, administrator login, and inventory management screens. The gender screen appears only after Formal Uniform. Include loading states, timeouts, retry behavior, and clear messages. The administrator uses a prototype PIN, follows the same conditional selection rule, reads the current quantity, enters a verified replacement quantity, and saves it. This PIN is for demonstration and must not be described as production security.",
        "",
        "Use a physical assistance button with INPUT_PULLUP, debounce it, sound a buzzer briefly, send one ASSIST|REQUESTED event, and show one local assistance message per press. Do not claim SMS, email, or internet notification.",
        "",
        "Use suggested Arduino pins D2 button, D6 green LED, D7 yellow LED, D8 red LED, D9 buzzer, D10 serial RX, and D11 serial TX unless the code and wiring documentation are updated together. The Uno uses 5 volt logic and the ESP32 uses 3.3 volt logic. Protect ESP32 RX from Uno TX using a verified level converter or voltage divider and connect a common ground.",
        "",
        "Implement the system in phases: hardware outputs, stock array, GET, SET, EEPROM, serial bridge, web server, student interface, administrator interface, assistance event, complete testing, then enclosure. Provide complete Arduino code, complete ESP32 code, complete web files, wiring documentation, setup instructions, and a functional test sheet. Clearly mark placeholders such as WiFi settings and the prototype PIN. Do not claim the system works until actual hardware tests are recorded.",
        "",
        "Keep the scope limited to checking and manually updating availability. Do not add payments, ordering, reservations, automatic deductions, barcode, RFID, official inventory integration, cloud synchronization, or a database unless the team requests a future version.",
        "```",
        "",
        "## Transfer Checklist",
        "",
        "- Give the next builder this brief and the complete Word manual.",
        "- Confirm that the builder repeats the no database requirement.",
        "- Confirm the 6 by 6 item and size order and the conditional gender rule before code is generated.",
        "- Confirm the exact stock thresholds and serial message formats.",
        "- Require code, wiring, setup steps, and tests as separate deliverables.",
        "- Replace sample quantities only after actual school stock is verified.",
        "- Record every approved change so the interface, firmware, wiring, and documentation remain consistent.",
        "",
        "## Defense Summary",
        "",
        "The prototype lets students check school uniform availability. The ESP32 receives browser requests and forwards them to the Arduino Uno. The Arduino reads the stored quantity, applies the stock rule, controls the matching LED, and returns the result. Authorized personnel manually replace verified quantities, and EEPROM preserves confirmed values after restart. The prototype has no database because only 36 values are required. It demonstrates hardware and software communication rather than a complete commercial inventory system.",
        "",
    ]
    return "\n".join(lines)


def append_transfer_sections(doc):
    doc.add_page_break()
    add_heading(doc, "Part Two Builder Transfer Specification", 1)
    add_para(doc, "This part is the implementation contract for another account, AI builder, programmer, or group member. It records the decisions that must remain consistent across the interface, ESP32 program, Arduino program, wiring, tests, and capstone defense.")

    add_heading(doc, "Transfer Package Purpose", 2)
    add_para(doc, "Use the Word manual for complete explanations and the companion Master Builder Brief for direct copy and paste into another account. The receiving builder should confirm the fixed decisions before generating code or changing the design.")

    add_heading(doc, "Fixed System Decisions", 2)
    add_bullets(doc, [
        "No external or cloud database is used. The Arduino stores 36 working quantities in a 6 by 6 array and preserves confirmed quantities in EEPROM.",
        "The Arduino Uno owns stock values, stock classification, EEPROM storage, LED control, assistance button reading, and buzzer control.",
        "The ESP32 serves the interface and bridges browser requests to the Arduino through serial communication.",
        "Student checks are read only. Only a validated administrator update may replace a quantity.",
        "Formal Uniform requires Male or Female. School T-Shirt and PE Uniform are unisex and skip gender.",
        "Available is 4 to 30, Low Stock is 1 to 3, and Out of Stock is 0.",
        "The prototype excludes purchases, payments, reservations, automatic deductions, barcode, RFID, official inventory integration, and production security.",
    ])

    add_heading(doc, "Implementation Responsibility Matrix", 2)
    add_table(doc, ["Part", "Required responsibility"], [
        ("HTML", "Create the screens, controls, forms, headings, and labels."),
        ("CSS", "Provide readable, consistent, and touchscreen friendly presentation."),
        ("JavaScript", "Control navigation, selections, requests, loading, validation, results, and errors."),
        ("ESP32", "Run the local web server, validate requests, translate HTTP to serial, and enforce timeouts."),
        ("Arduino Uno", "Own the inventory, apply status rules, save EEPROM values, and control physical devices."),
        ("Administrator", "Verify physical stock and enter accurate replacement quantities."),
    ], [1.35, 5.05])

    add_heading(doc, "Inventory and EEPROM Contract", 2)
    add_para(doc, "The item order, size order, and EEPROM formula must remain unchanged unless every related program, test, and document is updated together.")
    add_table(doc, ["Index", "Item code", "Display name", "EEPROM addresses"], [
        (0, "MALE_POLO", "Male Polo", "0 to 5"),
        (1, "FEMALE_BLOUSE", "Female Blouse", "6 to 11"),
        (2, "MALE_PANTS", "Male Pants", "12 to 17"),
        (3, "FEMALE_SKIRT", "Female Skirt", "18 to 23"),
        (4, "T_SHIRT", "School T-Shirt (Unisex)", "24 to 29"),
        (5, "PE_UNIFORM", "PE Uniform (Unisex)", "30 to 35"),
    ], [0.65, 1.65, 2.1, 1.8])
    add_para(doc, "Size order: XS, S, M, L, XL, XXL. Size indexes are 0 through 5. EEPROM address equals item index times 6 plus size index.", bold_lead="Size order:")

    add_heading(doc, "Gender and Unisex Selection Contract", 2)
    add_table(doc, ["Category", "Required choice", "Final stock item"], [
        ("Formal Uniform", "Male", "Male Polo or Male Pants"),
        ("Formal Uniform", "Female", "Female Blouse or Female Skirt"),
        ("School T-Shirt", "No gender", "School T-Shirt (Unisex)"),
        ("PE Uniform", "No gender", "PE Uniform (Unisex)"),
    ], [1.7, 1.6, 3.0])
    add_para(doc, "Gender controls only the Formal Uniform screen path. It is not another stored field or EEPROM dimension. The final item code and size are sufficient for stock lookup.")

    add_heading(doc, "Communication Contract", 2)
    add_code(doc, [
        "GET|T_SHIRT|M",
        "STOCK|12|AVAILABLE",
        "SET|T_SHIRT|M|15",
        "UPDATED|T_SHIRT|M|15",
        "ASSIST|REQUESTED",
        "ERROR|BAD_REQUEST",
        "ERROR|UNKNOWN_ITEM",
        "ERROR|UNKNOWN_SIZE",
        "ERROR|BAD_QUANTITY",
        "ERROR|TIMEOUT",
    ])
    add_para(doc, "Every command ends with a newline. The ESP32 must wait for an Arduino reply and must not display a successful update until UPDATED is received.")

    add_heading(doc, "Recommended Web Contract", 2)
    add_table(doc, ["Route", "Role", "Required behavior"], [
        ("GET /", "Open kiosk", "Return the student and administrator interface."),
        ("GET /check", "Read stock", "Validate item and size, send GET, and return a verified result."),
        ("POST /update", "Replace stock", "Require the prototype login, validate input, send SET, and wait for UPDATED."),
        ("GET /status", "Read local event", "Return assistance or controller state if this route is implemented."),
    ], [1.3, 1.5, 3.5])

    add_heading(doc, "Hardware Safety Contract", 2)
    add_table(doc, ["Function", "Suggested pin", "Condition"], [
        ("Assistance button", "D2", "Use INPUT_PULLUP and connect the button to ground."),
        ("Green LED", "D6", "Use a suitable current limiting resistor."),
        ("Yellow LED", "D7", "Use a suitable current limiting resistor."),
        ("Red LED", "D8", "Use a suitable current limiting resistor."),
        ("Buzzer", "D9", "Use a transistor driver if current requires it."),
        ("Serial RX", "D10", "Receive the ESP32 3.3 volt TX signal."),
        ("Serial TX", "D11", "Protect ESP32 RX from the Uno 5 volt output."),
    ], [2.1, 1.25, 2.95])
    add_para(doc, "Voltage safety. The Uno uses 5 volt logic and the ESP32 uses 3.3 volt logic. Use verified level conversion or a verified voltage divider from Uno TX to ESP32 RX, connect common ground, and remove power before rewiring.", bold_lead="Voltage safety.")

    add_heading(doc, "Required Build Order", 2)
    add_numbered(doc, [
        "Test the three LEDs, assistance button, and buzzer on the Arduino.",
        "Create the 6 by 6 stock array and verify all 36 positions through Serial Monitor.",
        "Implement GET and the exact stock thresholds.",
        "Implement SET validation in working memory before adding EEPROM.",
        "Add safe first run EEPROM initialization, loading, and single address updates.",
        "Establish ESP32 to Arduino serial communication using fixed test messages.",
        "Create the ESP32 web server and verify the check route from a browser.",
        "Build the student interface and connect it to real responses.",
        "Add the administrator login, current quantity lookup, validation, and confirmed update flow.",
        "Add button debounce, buzzer timing, and the local assistance message.",
        "Complete every functional test and record actual results and corrections.",
        "Build the enclosure only after the electronics and software work reliably on the bench.",
    ])

    add_heading(doc, "Builder Acceptance Checklist", 2)
    add_bullets(doc, [
        "All six stock item rows and all six sizes are selectable.",
        "Formal Uniform requires Male or Female and shows only the matching garments.",
        "School T-Shirt and PE Uniform are unisex and do not ask for gender.",
        "All 36 stock values can be checked.",
        "The exact Available, Low Stock, and Out of Stock thresholds are used everywhere.",
        "A student check never decreases a quantity.",
        "Only one status LED is on after a valid result.",
        "Administrator updates are validated twice, confirmed by Arduino, and retained after restart.",
        "Invalid updates keep the previous value.",
        "Disconnected hardware produces an error instead of a false result.",
        "One button press produces one buzzer event and one assistance message.",
        "The serial voltage level is protected and both boards share ground.",
        "Actual tests, dates, versions, results, and corrections are recorded.",
    ])

    add_heading(doc, "Instructions for Another Account", 2)
    add_para(doc, "Copy the complete Master Builder Brief into the other account before requesting code. Ask the receiving builder to repeat the fixed design, conditional gender rule, item order, size order, status thresholds, communication messages, and excluded features before it begins. Require complete Arduino code, ESP32 code, web files, wiring instructions, setup steps, and a test sheet as separate outputs.")

    add_heading(doc, "Change Control", 2)
    add_para(doc, "Do not change the storage method, pin assignments, item order, size order, thresholds, serial commands, or feature scope silently. If the group approves a change, update the Arduino program, ESP32 program, webpage, wiring diagram, tests, manual, and defense explanation together.")

    add_heading(doc, "Final Transfer Checklist", 2)
    add_bullets(doc, [
        "Transfer this complete Word manual and the Master Builder Brief.",
        "Confirm that no external database will be added to the current prototype.",
        "Confirm the exact 6 by 6 inventory mapping and conditional gender rule before generating code.",
        "Confirm the exact serial commands and status thresholds.",
        "Replace sample quantities only with verified values.",
        "Record actual hardware tests before claiming completion.",
        "Record every approved design change and keep all deliverables consistent.",
    ])


def main():
    runpy.run_path(str(ROOT / "build_formal_beginner_guide.py"), run_name="__main__")
    doc = Document(BASE)
    doc.paragraphs[0].text = "Arduino Uniform Stock Kiosk Complete Build and Transfer Manual"
    doc.paragraphs[1].text = "Grade 12 STEM Implementation Handout and Account Transfer Specification"
    doc.paragraphs[2].text = "School Uniform Stock Availability Prototype"
    for paragraph in doc.paragraphs:
        if paragraph.text == "Section 30 Sample Database Structure":
            paragraph.text = "Section 30 Sample Inventory Data Structure"
            paragraph.style = doc.styles["Heading 1"]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text == "Sample Database Structure":
                    cell.text = "Sample Inventory Data Structure"
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                paragraph.text = "Arduino Uniform Stock Kiosk Build and Transfer Manual"
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)
                    run.font.color.rgb = TEXT
    doc.core_properties.title = "Arduino Uniform Stock Kiosk Complete Build and Transfer Manual"
    doc.core_properties.subject = "Complete beginner friendly build manual and transferable implementation specification"
    doc.core_properties.author = "Grade 12 STEM Capstone Team"
    append_transfer_sections(doc)
    doc.save(OUT)
    BRIEF.write_text(build_markdown(), encoding="utf-8", newline="\n")
    print(OUT)
    print(BRIEF)


if __name__ == "__main__":
    main()
