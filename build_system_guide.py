from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output" / "documents"
OUT.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT / "Arduino_Uniform_Stock_Kiosk_System_Development_Guide.docx"

TITLE = "Development of an Arduino-Based Interactive Kiosk for School Uniform Stock Availability"
NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F3F6F9"
GRAY = "D9D9D9"
TEXT = "1F1F1F"
WHITE = "FFFFFF"


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRAY, size="6"):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_width(cell, width_inches):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(width_inches * 1440)))
    tcW.set(qn("w:type"), "dxa")


def set_keep_with_next(paragraph, keep=True):
    pPr = paragraph._p.get_or_add_pPr()
    el = pPr.find(qn("w:keepNext"))
    if keep and el is None:
        pPr.append(OxmlElement("w:keepNext"))
    elif not keep and el is not None:
        pPr.remove(el)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, display, end):
        run._r.append(el)


def set_font(run, name="Arial", size=None, bold=None, color=TEXT, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, bold_lead=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    if keep:
        set_keep_with_next(p)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        r = p.add_run(item)
        set_font(r)


def add_numbered(doc, items):
    for number, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{number}. {item}")
        set_font(r)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, size=15 if level == 1 else 14, bold=True, color="000000")
    set_keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.5, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_cant_split(hdr)
    for j, text in enumerate(headers):
        cell = hdr.cells[j]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[j])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(text))
        set_font(r, size=font_size, bold=True, color=WHITE)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        set_cant_split(table.rows[-1])
        for j, text in enumerate(row):
            cell = cells[j]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cell, widths[j])
            if i % 2 == 1:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if aligns and aligns[j] == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(text))
            set_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc, text, caption=None, size=8.5):
    if caption:
        p = doc.add_paragraph()
        r = p.add_run(caption)
        set_font(r, size=10, bold=True)
        set_keep_with_next(p)
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.1)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        set_font(r, name="Courier New", size=size, color="000000")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_diagram(doc, lines, caption=None):
    if caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption)
        set_font(r, size=10, bold=True)
        set_keep_with_next(p)
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(line)
        set_font(r, name="Courier New", size=9.5, bold=(line not in ("|", "v", "+")))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_warning(doc, lead, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(lead)
    set_font(r, bold=True, color="8B0000")
    r2 = p.add_run(" " + text)
    set_font(r2)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

# Normal and built-in styles
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    s = doc.styles[style_name]
    s.font.name = "Arial"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    s.font.size = Pt(11)
    s.paragraph_format.space_after = Pt(3)

title_style = doc.styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for n, size in ((1, 15), (2, 14), (3, 12)):
    s = doc.styles[f"Heading {n}"]
    s.font.name = "Arial"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(12 if n == 1 else 8)
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.keep_with_next = True

# Header and footer
hp = section.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hr = hp.add_run(TITLE)
set_font(hr, size=8, color="000000")
hp.paragraph_format.space_after = Pt(0)
fp = section.footer.paragraphs[0]
add_page_field(fp)
for r in fp.runs:
    set_font(r, size=9, color="4D4D4D")

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
p.paragraph_format.space_after = Pt(18)
r = p.add_run(TITLE)
set_font(r, size=24, bold=True, color="000000")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Complete System Development Guide")
set_font(r, size=16, bold=True, color=NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Grade 12 Capstone Prototype")
set_font(r, size=13, color="404040")
doc.add_paragraph()
for line in (
    "System focus  School uniform stock availability checking",
    "Core controller  Arduino Uno",
    "Communication bridge  ESP32",
    "User interface  HTML CSS and JavaScript kiosk",
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(line)
    set_font(r, size=11)
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("This guide explains what to build, how the parts communicate, what to test, and what remains outside the prototype scope.")
set_font(r, size=11, italic=True, color="404040")
doc.add_page_break()

add_heading(doc, "How to Use This Guide", 1)
add_para(doc, "This document is a practical reference for planning, building, testing, and demonstrating the prototype. Read Sections 1 to 4 first to understand the purpose and boundaries. Use Sections 5 to 16 while assembling the hardware and communication link. Use Sections 17 to 30 while building the interface and programs. Complete the tests in Sections 31 to 34 before the capstone defense.")
add_warning(doc, "Prototype status.", "The diagrams, sample quantities, commands, pin assignments, and programs are starting points. Confirm the final wiring and record actual test results from the assembled kiosk before claiming that the system works.")

contents = [
    (1, "Project Overview"), (2, "Project Objectives"), (3, "Scope of the System"),
    (4, "Project Limitations"), (5, "Complete Hardware Requirements"), (6, "Role of the Arduino Uno"),
    (7, "Role of the ESP32"), (8, "Software Requirements"), (9, "Complete System Architecture"),
    (10, "How the System Works"), (11, "Stock Classification Logic"), (12, "Sample School Uniform Inventory"),
    (13, "Data Storage"), (14, "Serial Communication"), (15, "Important Voltage Safety"),
    (16, "Hardware Wiring"), (17, "Kiosk User Interface"), (18, "Admin Interface"),
    (19, "Admin Communication Command"), (20, "Assistance Button"), (21, "LED Indicator Panel"),
    (22, "Proposed Physical Kiosk Design"), (23, "Complete User Flowchart"), (24, "Admin Flowchart"),
    (25, "Example Arduino Program Structure"), (26, "Example ESP32 Program Structure"), (27, "HTML Structure"),
    (28, "CSS Design"), (29, "JavaScript Function"), (30, "Sample Database Structure"),
    (31, "Functional Testing"), (32, "Possible Problems and Troubleshooting"), (33, "Development Order"),
    (34, "Demonstration Scenario"), (35, "Questions the Panel May Ask"), (36, "Future Improvements"),
    (37, "Final Recommended System"),
]
add_heading(doc, "Contents", 2)
add_table(doc, ["Section", "Topic"], contents, widths=[0.8, 5.9], font_size=9.2, aligns=["center", "left"])
doc.add_page_break()

# Section 1
add_heading(doc, "Section 1 Project Overview", 1)
add_para(doc, "The proposed system is a self-service kiosk that lets students check whether a required school uniform and size are currently in stock. A student uses a touchscreen or computer display, chooses an item and size, and receives a quantity and a simple status: Available, Low Stock, or Out of Stock. A matching LED provides a physical signal beside the screen.")
add_para(doc, "The kiosk may be placed near the uniform office, cashier area, student services office, or another supervised school location. It addresses a simple problem: students may wait in line or ask staff about an item that is unavailable. A quick stock check can reduce repeated questions and help a student decide whether to proceed to the uniform office.")
add_para(doc, "An interactive kiosk is appropriate because it presents only the choices needed for the task. Large buttons make the process easy to follow, while the physical LEDs and assistance button show that the webpage and electronics work together.")
add_para(doc, "The Arduino Uno performs the stock and hardware work. It reads stock values, classifies availability, controls the LEDs and buzzer, and reads the assistance button. The ESP32 provides the network and webpage connection. It receives requests from the browser and forwards them to the Arduino Uno through serial communication.")
add_warning(doc, "Main focus.", "This prototype checks stock availability. It does not sell uniforms, accept payment, reserve items, dispense products, or automatically change stock after a student checks an item.")

# Section 2
add_heading(doc, "Section 2 Project Objectives", 1)
add_heading(doc, "General Objective", 2)
add_para(doc, "To develop an Arduino-based interactive kiosk prototype that allows students to check the quantity and availability status of school uniforms while providing physical status indicators and a simple administrator method for maintaining stock values.")
add_heading(doc, "Specific Objectives", 2)
add_bullets(doc, [
    "Develop a touchscreen-friendly kiosk interface using HTML, CSS, and JavaScript.",
    "Allow students to browse the supported uniform types and select a size.",
    "Retrieve and display the current stock quantity for the selected item and size.",
    "Classify stock as Available, Low Stock, or Out of Stock using clear prototype rules.",
    "Control green, yellow, and red LEDs so the physical indicator matches the displayed status.",
    "Provide a physical assistance button that activates a local buzzer and displays an assistance message.",
    "Allow an authorized administrator to view and manually update stock quantities through a simple PIN-protected screen.",
    "Store prototype stock values so they can be restored after the Arduino restarts.",
    "Demonstrate reliable communication among the kiosk browser, ESP32, Arduino Uno, and connected hardware.",
])

# Section 3
add_heading(doc, "Section 3 Scope of the System", 1)
add_para(doc, "The scope defines the functions that must work in the Grade 12 prototype. A feature outside this list should not be added unless the group has already completed and tested the required functions.")
add_heading(doc, "Student Functions", 2)
add_bullets(doc, [
    "Open the kiosk home screen.", "Browse the supported uniform types.", "Select one uniform and one size.",
    "Send an availability request.", "View the stored stock quantity and availability status.",
    "See the corresponding LED turn on.", "Request local assistance.",
    "Check another item or return to the home screen.",
])
add_heading(doc, "Administrator Functions", 2)
add_bullets(doc, [
    "Open the administrator login screen and enter a prototype PIN.",
    "Select a uniform and size and view its current stored quantity.",
    "Enter a non-negative replacement quantity.",
    "Send the update to the Arduino Uno and save it to EEPROM.",
    "Receive a confirmation showing the item, size, and new quantity.",
])
add_heading(doc, "Hardware Functions", 2)
add_table(doc, ["Hardware", "Required behavior"], [
    ("Green LED", "Turns on for Available when stock is 4 to 30."),
    ("Yellow LED", "Turns on for Low Stock when stock is 1 to 3."),
    ("Red LED", "Turns on for Out of Stock when stock is 0."),
    ("Buzzer", "Sounds briefly for confirmation or assistance. It must not remain on continuously."),
    ("Assistance button", "Sends a local assistance event when pressed."),
], widths=[1.4, 5.3], aligns=["center", "left"])

# Section 4
add_heading(doc, "Section 4 Project Limitations", 1)
add_para(doc, "The prototype does not include actual purchasing, online ordering, payment processing, GCash, credit or debit cards, reservations, automatic inventory deduction, RFID tracking, automatic dispensing, student accounts, facial recognition, an AI chatbot, SMS notifications, or cloud-based enterprise inventory management.")
add_para(doc, "An authorized administrator manually maintains the stock information. The kiosk reports the number stored in the Arduino; it does not independently confirm the quantity physically present in the uniform office. Staff must count or verify the actual items and enter the correct number.")
add_warning(doc, "Checking does not subtract stock.", "If PE Shirt Medium contains 12 pieces, a student check must still leave the value at 12. Only an administrator SET command changes the stored quantity.")
add_para(doc, "The PIN is suitable only for a supervised prototype. A real deployment would require stronger authentication, protected network access, audit records, backup procedures, and integration with the school's official inventory process.")
# Section 5
add_heading(doc, "Section 5 Complete Hardware Requirements", 1)
hardware_rows = [
    ("Arduino Uno R3", "1", "Runs stock logic and controls local hardware.", "Main controller; use an original or compatible board."),
    ("ESP32 DevKit", "1", "Hosts or connects the kiosk webpage and bridges requests to Arduino.", "Uses 3.3 V GPIO; do not feed 5 V into its pins."),
    ("Laptop mini PC or computer", "1", "Displays the kiosk webpage.", "A laptop is acceptable for the prototype."),
    ("Touchscreen monitor", "Optional 1", "Provides direct touch input.", "A mouse may be used during early testing."),
    ("Green LED", "1", "Shows Available status.", "Use with a current-limiting resistor."),
    ("Yellow LED", "1", "Shows Low Stock status.", "Use with a current-limiting resistor."),
    ("Red LED", "1", "Shows Out of Stock status.", "Use with a current-limiting resistor."),
    ("220 ohm resistors", "3", "Limit current through the three LEDs.", "One resistor per LED."),
    ("Active buzzer", "1", "Produces a simple notification tone.", "Check its voltage and current rating."),
    ("Push button", "1", "Lets a student request assistance.", "Use Arduino INPUT_PULLUP wiring."),
    ("2N2222 transistor", "1 if needed", "Lets Arduino switch a buzzer that needs more current.", "Use when buzzer current is above safe GPIO current."),
    ("1 kilohm resistor", "1", "Limits current into the transistor base.", "Used with the transistor driver."),
    ("Breadboard", "1", "Holds prototype wiring without soldering.", "Use labeled rails and tidy wire routes."),
    ("Jumper wires", "As needed", "Connect boards, LEDs, button, and buzzer.", "Use male to male or male to female as required."),
    ("USB cables", "2", "Power and program the Arduino and ESP32.", "Use data-capable cables, not charge-only cables."),
    ("Logic-level converter", "Optional 1", "Safely translates Uno 5 V TX to ESP32 3.3 V RX.", "Recommended for a reliable serial link."),
    ("Kiosk enclosure", "1", "Holds the display and protects electronics.", "Provide ventilation and service access."),
]
add_table(doc, ["Component", "Quantity", "Purpose", "Important Notes"], hardware_rows, widths=[1.65, 0.85, 2.15, 2.05], font_size=8.5, aligns=["left", "center", "left", "left"])
add_warning(doc, "Power note.", "Do not power a laptop, monitor, or high-current buzzer from an Arduino pin. Use the proper rated power source for each device and connect grounds only as required by the communication design.")

# Section 6
add_heading(doc, "Section 6 Role of the Arduino Uno", 1)
add_para(doc, "The Arduino Uno is the hardware controller and the owner of the prototype stock logic. It receives a request such as GET|PE_SHIRT|M, finds the stored quantity, applies the availability rule, activates the correct LED, and returns a response to the ESP32.")
add_bullets(doc, [
    "Store or access the 36 prototype stock values for six uniforms and six sizes.",
    "Process GET commands for stock checks and SET commands for administrator updates.",
    "Classify the quantity as Available, Low Stock, or Out of Stock.",
    "Ensure only one stock-status LED is active after a check.",
    "Control the buzzer and read the physical assistance button.",
    "Load saved values from EEPROM at startup and update EEPROM only when a value changes.",
    "Exchange newline-terminated text commands with the ESP32.",
])
add_para(doc, "The Arduino is therefore not decorative hardware. The demonstration should show that disconnecting or stopping the Arduino prevents stock requests and physical indicators from working. This proves that the Arduino performs essential processing.")

# Section 7
add_heading(doc, "Section 7 Role of the ESP32", 1)
add_para(doc, "The ESP32 is the communication bridge. It may create a local Wi-Fi access point or join a school-approved local network. It serves the kiosk webpage or responds to browser requests, converts web requests into simple serial commands, waits for the Arduino response, and returns the result to JavaScript.")
add_table(doc, ["Device", "Main responsibility"], [
    ("Arduino Uno", "Hardware controller and stock logic"),
    ("ESP32", "Communication bridge and web server"),
    ("Computer or touchscreen", "Graphical user interface used by the student or administrator"),
], widths=[2.0, 4.7], aligns=["center", "left"])
add_para(doc, "Using both boards makes each responsibility easy to explain. The ESP32 is better suited to Wi-Fi and web requests, while the Arduino Uno gives the project a clear physical-computing controller for stock rules, LEDs, buzzer, and button.")

# Section 8
add_heading(doc, "Section 8 Software Requirements", 1)
software_rows = [
    ("Arduino IDE", "Writes, compiles, and uploads programs to the Arduino Uno and ESP32."),
    ("HTML", "Defines the text, buttons, forms, and sections shown on the kiosk screen."),
    ("CSS", "Controls sizes, spacing, colors, and touchscreen-friendly layout."),
    ("JavaScript", "Changes screens, stores selections, sends requests, and displays results."),
    ("Chrome or Edge", "Opens the local kiosk webpage in full-screen or kiosk mode."),
    ("ESP32 WebServer library", "Creates routes such as /check, /update, /status, and /assist."),
    ("Arduino EEPROM library", "Saves small stock values so the Uno can restore them after power loss."),
]
add_table(doc, ["Software", "Purpose"], software_rows, widths=[2.0, 4.7], aligns=["left", "left"])
add_para(doc, "React, Node.js, Firebase, MySQL, Supabase, and similar platforms are not required for this basic prototype. Plain HTML, CSS, JavaScript, the ESP32 WebServer library, and Arduino EEPROM are sufficient. Additional technology should be considered only after the core system works reliably.")
# Section 9
add_heading(doc, "Section 9 Complete System Architecture", 1)
add_diagram(doc, [
    "STUDENT", "|", "v", "TOUCHSCREEN OR COMPUTER", "|", "v",
    "HTML + CSS + JAVASCRIPT KIOSK", "|  HTTP over local Wi-Fi", "v",
    "ESP32 WEB SERVER", "|  UART serial commands", "v", "ARDUINO UNO",
    "|", "v", "STOCK DATA + AVAILABILITY LOGIC", "|", "v",
    "GREEN LED + YELLOW LED + RED LED + BUZZER + ASSISTANCE BUTTON",
], "System block diagram")
architecture_rows = [
    ("Student", "Chooses an item, selects a size, reads the result, or requests assistance."),
    ("Touchscreen or computer", "Displays the browser and sends touch or mouse input."),
    ("HTML CSS JavaScript", "Presents screens, validates choices, and calls ESP32 routes."),
    ("ESP32 Web Server", "Receives HTTP requests and translates them into serial messages."),
    ("UART serial communication", "Carries short text commands between the two boards."),
    ("Arduino Uno", "Looks up stock, applies rules, saves updates, and controls hardware."),
    ("Indicators and controls", "Show the result and provide local assistance input."),
]
add_table(doc, ["Stage", "What happens"], architecture_rows, widths=[2.0, 4.7], aligns=["left", "left"])
add_para(doc, "A request travels down the diagram, while the result travels back up. For example, JavaScript asks the ESP32 to check PE Shirt Medium. The ESP32 sends GET|PE_SHIRT|M to the Arduino. The Arduino replies STOCK|12|AVAILABLE, and the ESP32 returns that information to the browser.")

# Section 10
add_heading(doc, "Section 10 How the System Works", 1)
add_numbered(doc, [
    "The student approaches the kiosk and sees the home screen.",
    "The student presses Check Availability.",
    "The interface displays the uniform categories.",
    "The student selects one uniform type.",
    "The interface displays the supported sizes.",
    "The student selects a size and reviews the confirmation screen.",
    "The student presses Check Availability.",
    "JavaScript sends an HTTP request to the ESP32, such as /check?item=PE_SHIRT&size=M.",
    "The ESP32 validates the parameters and sends GET|PE_SHIRT|M followed by a newline to the Arduino Uno.",
    "The Arduino searches its stock array for the matching item and size.",
    "The Arduino determines whether the quantity is Available, Low Stock, or Out of Stock.",
    "The Arduino turns off all status LEDs and then turns on only the correct LED.",
    "The Arduino sends a response such as STOCK|12|AVAILABLE to the ESP32.",
    "The ESP32 returns a small response to the browser.",
    "JavaScript displays the quantity, status, guidance, and Check Another and Home buttons.",
    "The student checks another item or returns to the home screen.",
])
add_para(doc, "If the Arduino does not respond within the ESP32 timeout, the browser should show a clear error such as Hardware controller not responding. It must not invent a stock value.")

# Section 11
add_heading(doc, "Section 11 Stock Classification Logic", 1)
add_table(doc, ["Stored quantity", "Status", "LED"], [
    ("4 to 30", "AVAILABLE", "Green"),
    ("1 to 3", "LOW STOCK", "Yellow"),
    ("0", "OUT OF STOCK", "Red"),
], widths=[2.0, 2.5, 2.2], aligns=["center", "center", "center"])
add_code(doc, """IF stock >= 4
    status = AVAILABLE
ELSE IF stock >= 1
    status = LOW STOCK
ELSE
    status = OUT OF STOCK
END IF""", "Availability pseudocode")
add_para(doc, "These thresholds are prototype rules. The school may later change them based on actual demand and restocking practice. The same thresholds must be used by the Arduino, interface labels, tests, and defense explanation.")

# Section 12
add_heading(doc, "Section 12 Sample School Uniform Inventory", 1)
add_warning(doc, "Sample data only.", "Replace these quantities with approved demonstration values or verified school inventory. Do not present this table as the school's actual stock record.")
inventory = [
    ("Male Polo", 12, 8, 5, 3, 2, 0),
    ("Female Blouse", 9, 7, 6, 4, 1, 0),
    ("Male Pants", 6, 5, 4, 3, 2, 1),
    ("Female Skirt", 8, 7, 5, 4, 2, 0),
    ("PE Shirt", 10, 15, 12, 8, 4, 2),
    ("PE Pants", 7, 9, 6, 3, 1, 0),
]
add_table(doc, ["Uniform", "XS", "S", "M", "L", "XL", "XXL"], inventory, widths=[1.7, .75, .75, .75, .75, .75, .75], font_size=9, aligns=["left", "center", "center", "center", "center", "center", "center"])
add_para(doc, "The table intentionally includes Available, Low Stock, and Out of Stock examples. This makes it useful for testing all three LED and result states.")

# Section 13
add_heading(doc, "Section 13 Data Storage", 1)
add_para(doc, "EEPROM is a small memory area that can keep data after power is removed. The Arduino Uno can use it to remember stock values after a restart. In this prototype, each uniform and size combination can be assigned an address.")
add_para(doc, "Six uniform types multiplied by six sizes produces 36 stock values. Each demonstration quantity is limited to 0 through 30, and one byte stores each value. Address 0 may represent Male Polo XS, address 1 Male Polo S, and so on. A formula such as itemIndex multiplied by 6 plus sizeIndex gives the address.")
add_code(doc, """address = (itemIndex * 6) + sizeIndex
quantity = EEPROM.read(address)

IF newQuantity is different from quantity
    EEPROM.update(address, newQuantity)
END IF""", "Conceptual EEPROM mapping")
add_warning(doc, "EEPROM write limit.", "EEPROM can be written only a limited number of times. Use EEPROM.update and save only when an administrator confirms a changed value. Do not rewrite all values during every loop cycle.")
add_para(doc, "A first-run marker can tell the program whether valid values already exist. If the marker is missing, the Arduino writes the sample defaults once. On later starts, it reads the saved quantities instead of replacing them.")

# Section 14
add_heading(doc, "Section 14 Serial Communication", 1)
add_para(doc, "Serial communication sends characters one after another between the ESP32 and Arduino Uno. UART is the hardware method used for this link. Both programs should use the same baud rate, such as 9600 or 115200, and end each command with a newline character so the receiver knows the message is complete.")
add_table(doc, ["Direction", "Example", "Meaning"], [
    ("ESP32 to Arduino", "GET|PE_SHIRT|M", "Request the Medium PE Shirt quantity."),
    ("Arduino to ESP32", "STOCK|12|AVAILABLE", "Return quantity 12 with Available status."),
    ("ESP32 to Arduino", "GET|POLO|XL", "Request the Extra Large polo quantity."),
    ("Arduino to ESP32", "STOCK|2|LOW", "Return quantity 2 with Low Stock status."),
    ("ESP32 to Arduino", "GET|POLO|XXL", "Request the Double Extra Large polo quantity."),
    ("Arduino to ESP32", "STOCK|0|OUT", "Return quantity 0 with Out of Stock status."),
], widths=[1.45, 2.15, 3.1], font_size=9, aligns=["center", "center", "left"])
add_para(doc, "The vertical bar separates fields. In GET|PE_SHIRT|M, GET is the action, PE_SHIRT is the item code, and M is the size code. The programs must use exactly the same item and size codes. Unknown or incomplete commands should return an error such as ERROR|BAD_REQUEST.")

# Section 15
add_heading(doc, "Section 15 Important Voltage Safety", 1)
add_para(doc, "Arduino Uno digital pins normally use 5 V logic. ESP32 GPIO pins use approximately 3.3 V logic and are not designed to receive 5 V. A direct connection from an Arduino TX pin to an ESP32 RX pin can damage the ESP32.")
add_warning(doc, "Do not connect a 5 V Uno TX output directly to an ESP32 GPIO.", "Use a proper bidirectional logic-level converter or a correctly calculated resistor voltage divider on the Uno-to-ESP32 signal. Ask a teacher or electronics adviser to verify the circuit before power is applied.")
add_para(doc, "The ESP32 TX output is 3.3 V. It is usually read as HIGH by the Arduino Uno, but the final circuit should still be checked. The Arduino and ESP32 must share a common GND for the serial voltage levels to have the same reference.")
add_para(doc, "Connect TX to the other device's RX, not TX to TX. Keep the serial wires short during breadboard testing. Disconnect power before changing wiring.")

# Section 16
add_heading(doc, "Section 16 Hardware Wiring", 1)
add_table(doc, ["Function", "Suggested connection"], [
    ("Green LED", "Arduino D6 -> 220 ohm resistor -> green LED anode; LED cathode -> GND"),
    ("Yellow LED", "Arduino D7 -> 220 ohm resistor -> yellow LED anode; LED cathode -> GND"),
    ("Red LED", "Arduino D8 -> 220 ohm resistor -> red LED anode; LED cathode -> GND"),
    ("Buzzer", "Arduino D9 -> 1 kilohm resistor -> transistor base; transistor switches buzzer as rated"),
    ("Assistance button", "Arduino D2 -> push button -> GND; configure D2 as INPUT_PULLUP"),
    ("ESP32 TX2", "ESP32 GPIO17 TX -> Arduino D10 SoftwareSerial RX"),
    ("Arduino serial TX", "Arduino D11 SoftwareSerial TX -> level shifter or divider -> ESP32 GPIO16 RX2"),
    ("Common reference", "Arduino GND -> ESP32 GND"),
], widths=[1.7, 5.0], font_size=9.2, aligns=["left", "left"])
add_para(doc, "With INPUT_PULLUP, the Arduino enables an internal resistor. The button does not need a separate pull-up resistor in the basic circuit. Button not pressed reads HIGH. Button pressed connects the pin to ground and reads LOW.")
add_para(doc, "The suggested D10 and D11 pins use SoftwareSerial so the Uno USB serial port remains available for programming and debugging. If different pins are used, update the wiring diagram and both programs.")
add_heading(doc, "Transistor Buzzer Connection", 2)
add_bullets(doc, [
    "Connect Arduino D9 through a 1 kilohm resistor to the transistor base.",
    "Connect the transistor emitter to GND.",
    "Connect the buzzer negative terminal to the transistor collector.",
    "Connect the buzzer positive terminal to its rated supply voltage.",
    "Connect the supply ground to Arduino GND.",
    "For an inductive sounder, add the protection component recommended by its datasheet or adviser.",
])
# Section 17
add_heading(doc, "Section 17 Kiosk User Interface", 1)
add_para(doc, "The interface should use large touch targets, short instructions, strong contrast, and a consistent Home and Back location. The student should always know what was selected and what to do next.")
screens = [
    ("Screen 1 Home", "Title: School Uniform Availability Kiosk. Buttons: Check Availability, Request Assistance, Uniform Information."),
    ("Screen 2 Uniform Selection", "Buttons: Male Polo, Female Blouse, Male Pants, Female Skirt, PE Shirt, PE Pants, and Back."),
    ("Screen 3 Size Selection", "Buttons: XS, S, M, L, XL, XXL, and Back. Show the selected uniform above the choices."),
    ("Screen 4 Confirmation", "Show Selected Item and Selected Size. Buttons: Check Availability and Back."),
    ("Screen 5 Result", "Show a large status, item, size, quantity, guidance, and buttons for Check Another and Home."),
]
add_table(doc, ["Screen", "Required content"], screens, widths=[1.65, 5.05], font_size=9.2, aligns=["left", "left"])
add_heading(doc, "Result Wording", 2)
add_table(doc, ["Status", "Example screen message"], [
    ("AVAILABLE", "PE Shirt | Medium | 12 Pieces Available | Please proceed to the Uniform Office for purchasing."),
    ("LOW STOCK", "Only 2 pieces remaining. Please verify availability with the Uniform Office."),
    ("OUT OF STOCK", "This size is currently unavailable. Please check another size or ask school personnel."),
], widths=[1.55, 5.15], aligns=["center", "left"])
add_para(doc, "Do not use green, yellow, or red as the only way to communicate status. Always display the status word and quantity because color can be misunderstood and the physical LED may not be visible from every angle.")

# Section 18
add_heading(doc, "Section 18 Admin Interface", 1)
add_para(doc, "The administrator area is separate from the student screens. It begins with a simple prototype login containing a PIN field and Login button. After a correct PIN, the Inventory Management screen shows a uniform dropdown, size dropdown, current stock, new quantity input, Update Stock button, and Logout button.")
add_numbered(doc, [
    "The administrator chooses a uniform and size.",
    "JavaScript requests the current value through the normal check route.",
    "The Arduino returns the existing quantity.",
    "The administrator enters a whole number of zero or greater.",
    "The ESP32 sends a SET command to the Arduino.",
    "The Arduino validates the command, changes the array value, and saves only that address to EEPROM.",
    "The Arduino returns an UPDATED response, and the interface displays a confirmation.",
])
add_warning(doc, "Validate every update.", "Reject blank input, negative numbers, decimals, values above the selected storage limit, unknown items, and unknown sizes. Do not display success until the Arduino confirms the saved value.")
add_para(doc, "For a classroom prototype, the PIN may be stored in the ESP32 program. Explain during the defense that this is not secure enough for actual school deployment because code and network traffic may be inspected.")

# Section 19
add_heading(doc, "Section 19 Admin Communication Command", 1)
add_code(doc, """SET|PE_SHIRT|M|15
UPDATED|PE_SHIRT|M|15""", "Example administrator update")
add_para(doc, "SET tells the Arduino to replace a value. PE_SHIRT identifies the uniform, M identifies the size, and 15 is the new quantity. The response confirms that Medium PE Shirt stock is now 15. A later GET request should return STOCK|15|AVAILABLE.")
add_para(doc, "If the request is invalid, the Arduino should keep the old value and return a message such as ERROR|BAD_QUANTITY or ERROR|UNKNOWN_ITEM. This prevents a failed update from appearing successful.")

# Section 20
add_heading(doc, "Section 20 Assistance Button", 1)
add_numbered(doc, [
    "The student presses the physical assistance button.",
    "The Arduino detects a change from HIGH to LOW.",
    "The Arduino sounds the buzzer for a short, controlled period.",
    "The Arduino sends ASSIST|REQUESTED to the ESP32.",
    "The ESP32 makes the event available to the webpage through /status or /assist.",
    "The kiosk displays Assistance Requested. Please wait for school personnel.",
])
add_para(doc, "The program should debounce the button so one press creates only one event. A delay of about 30 to 50 milliseconds or a non-blocking time check can filter electrical bouncing.")
add_warning(doc, "Local notification only.", "The prototype buzzer and screen message do not automatically contact a person through SMS, email, or the internet. School personnel must be near enough to hear or see the request.")

# Section 21
add_heading(doc, "Section 21 LED Indicator Panel", 1)
add_diagram(doc, [
    "STOCK STATUS", "GREEN   Available", "YELLOW  Low Stock", "RED     Out of Stock"
], "Suggested label beside or below the screen")
add_para(doc, "Mount the LEDs where students and panel members can see them. Label each color with its meaning. After a check, the Arduino should first switch all three LEDs off and then switch on the one that matches the result. When the user returns home, the program may turn all status LEDs off so an old result is not mistaken for the current one.")

# Section 22
add_heading(doc, "Section 22 Proposed Physical Kiosk Design", 1)
add_para(doc, "A simple enclosure may be built from plywood, acrylic, PVC board, or another safe prototype material. Place the touchscreen or monitor near eye level. Put the three labeled LEDs and the assistance button below or beside the screen. Keep the buzzer behind a small opening so it can be heard without being touched.")
add_bullets(doc, [
    "Mount the Arduino Uno and ESP32 inside the enclosure on non-conductive standoffs.",
    "Separate loose wires from moving parts and screen brackets.",
    "Provide ventilation openings for the computer and power components.",
    "Keep USB and power connectors accessible for programming and maintenance.",
    "Use strain relief so a cable cannot pull directly on a circuit board.",
    "Cover exposed conductors and avoid sharp edges.",
    "Add labels for Green Available, Yellow Low Stock, Red Out of Stock, and Request Assistance.",
])
add_para(doc, "For the defense, a laptop on a stable stand with a small electronics panel is acceptable if a full enclosure is not yet safe or complete. The important requirement is a working and clearly explained system, not an expensive cabinet.")
# Section 23
add_heading(doc, "Section 23 Complete User Flowchart", 1)
add_diagram(doc, [
    "START", "|", "v", "HOME SCREEN", "|", "v", "CHECK AVAILABILITY", "|", "v",
    "SELECT UNIFORM", "|", "v", "SELECT SIZE", "|", "v", "CONFIRM", "|", "v",
    "SEND REQUEST TO ESP32", "|", "v", "ESP32 SENDS GET TO ARDUINO", "|", "v",
    "ARDUINO READS STOCK AND CLASSIFIES IT", "|", "v",
    "AVAILABLE -> GREEN   LOW -> YELLOW   NONE -> RED", "|", "v",
    "SEND RESULT TO ESP32", "|", "v", "DISPLAY RESULT", "|", "v",
    "CHECK ANOTHER", "| YES                         | NO", "v                             v",
    "UNIFORM SELECTION              HOME SCREEN",
], "Student availability flow")
add_para(doc, "A timeout or invalid response should branch to an error screen with Retry and Home. It should not enter the Available, Low Stock, or Out of Stock branches.")

# Section 24
add_heading(doc, "Section 24 Admin Flowchart", 1)
add_diagram(doc, [
    "START", "|", "v", "ADMIN LOGIN", "|", "v", "ENTER PIN", "|", "v",
    "PIN CORRECT", "| NO -> SHOW ERROR AND RETURN TO LOGIN", "| YES", "v",
    "INVENTORY MANAGEMENT", "|", "v", "SELECT UNIFORM AND SIZE", "|", "v",
    "DISPLAY CURRENT QUANTITY", "|", "v", "ENTER NEW QUANTITY", "|", "v",
    "VALID INPUT", "| NO -> SHOW VALIDATION MESSAGE", "| YES", "v", "SEND SET COMMAND", "|", "v",
    "ARDUINO SAVES VALUE", "|", "v", "EEPROM UPDATED", "|", "v", "SHOW CONFIRMATION",
], "Administrator update flow")

# Section 25
add_heading(doc, "Section 25 Example Arduino Program Structure", 1)
add_para(doc, "The Uno program should stay small and predictable. Separate item lookup, size lookup, stock classification, LED control, command processing, and button monitoring into simple functions. The code below is a beginner-friendly structure, not a finished production program.")
arduino_code = r'''#include <EEPROM.h>
#include <SoftwareSerial.h>

const byte GREEN_LED = 6;
const byte YELLOW_LED = 7;
const byte RED_LED = 8;
const byte BUZZER = 9;
const byte ASSIST_BUTTON = 2;

SoftwareSerial espSerial(10, 11); // RX, TX

const byte ITEM_COUNT = 6;
const byte SIZE_COUNT = 6;
byte stock[ITEM_COUNT][SIZE_COUNT];

const char* itemCodes[ITEM_COUNT] = {
  "POLO", "BLOUSE", "MALE_PANTS",
  "FEMALE_SKIRT", "PE_SHIRT", "PE_PANTS"
};
const char* sizeCodes[SIZE_COUNT] = {"XS", "S", "M", "L", "XL", "XXL"};

void setup() {
  pinMode(GREEN_LED, OUTPUT);
  pinMode(YELLOW_LED, OUTPUT);
  pinMode(RED_LED, OUTPUT);
  pinMode(BUZZER, OUTPUT);
  pinMode(ASSIST_BUTTON, INPUT_PULLUP);

  Serial.begin(9600);       // USB debugging
  espSerial.begin(9600);    // link to ESP32
  loadInventory();
  setStatusLeds(-1);
}

void loop() {
  if (espSerial.available()) {
    String command = espSerial.readStringUntil('\n');
    command.trim();
    processCommand(command);
  }
  monitorAssistanceButton();
}

void processCommand(String command) {
  if (command.startsWith("GET|")) {
    processGet(command);
  } else if (command.startsWith("SET|")) {
    processSet(command);
  } else {
    espSerial.println("ERROR|BAD_REQUEST");
  }
}

String statusFor(byte quantity) {
  if (quantity >= 4) return "AVAILABLE";
  if (quantity >= 1) return "LOW";
  return "OUT";
}

void setStatusLeds(int quantity) {
  digitalWrite(GREEN_LED, LOW);
  digitalWrite(YELLOW_LED, LOW);
  digitalWrite(RED_LED, LOW);
  if (quantity < 0) return;
  if (quantity >= 4) digitalWrite(GREEN_LED, HIGH);
  else if (quantity >= 1) digitalWrite(YELLOW_LED, HIGH);
  else digitalWrite(RED_LED, HIGH);
}

void processGet(String command) {
  String item = field(command, 1);
  String size = field(command, 2);
  int itemIndex = findCode(itemCodes, ITEM_COUNT, item);
  int sizeIndex = findCode(sizeCodes, SIZE_COUNT, size);
  if (itemIndex < 0 || sizeIndex < 0) {
    espSerial.println("ERROR|UNKNOWN_ITEM_OR_SIZE");
    return;
  }
  byte quantity = stock[itemIndex][sizeIndex];
  setStatusLeds(quantity);
  espSerial.println("STOCK|" + String(quantity) + "|" + statusFor(quantity));
}

void processSet(String command) {
  String item = field(command, 1);
  String size = field(command, 2);
  String valueText = field(command, 3);
  int itemIndex = findCode(itemCodes, ITEM_COUNT, item);
  int sizeIndex = findCode(sizeCodes, SIZE_COUNT, size);
  int value = valueText.toInt();
  if (itemIndex < 0 || sizeIndex < 0 || value < 0 || value > 30) {
    espSerial.println("ERROR|BAD_UPDATE");
    return;
  }
  stock[itemIndex][sizeIndex] = (byte)value;
  int address = itemIndex * SIZE_COUNT + sizeIndex;
  EEPROM.update(address, (byte)value);
  espSerial.println("UPDATED|" + item + "|" + size + "|" + String(value));
}

void monitorAssistanceButton() {
  static bool wasPressed = false;
  bool pressed = digitalRead(ASSIST_BUTTON) == LOW;
  if (pressed && !wasPressed) {
    digitalWrite(BUZZER, HIGH);
    delay(150);
    digitalWrite(BUZZER, LOW);
    espSerial.println("ASSIST|REQUESTED");
  }
  wasPressed = pressed;
  delay(20);
}

void loadInventory() {
  for (byte item = 0; item < ITEM_COUNT; item++) {
    for (byte size = 0; size < SIZE_COUNT; size++) {
      stock[item][size] = EEPROM.read(item * SIZE_COUNT + size);
    }
  }
}

// Implement field() to return text between | separators.
// Implement findCode() to return a matching array index or -1.'''
add_code(doc, arduino_code, "Simplified Arduino Uno program structure", size=7.7)
add_para(doc, "Before using the code, add a safe first-run initialization routine. A new Uno EEPROM may contain 255, which would appear as a large stock value. Use a marker byte outside the 36 stock addresses. If the marker is absent, write the sample data once and then write the marker.")

# Section 26
add_heading(doc, "Section 26 Example ESP32 Program Structure", 1)
add_para(doc, "The ESP32 program connects the browser to the Uno. For a demonstration that does not depend on school Wi-Fi, the ESP32 can create its own local access point. The kiosk computer connects to that network and opens the ESP32 address.")
esp_code = r'''#include <WiFi.h>
#include <WebServer.h>

WebServer server(80);
HardwareSerial unoSerial(2);

String askUno(String command) {
  while (unoSerial.available()) unoSerial.read();
  unoSerial.println(command);
  unsigned long start = millis();
  while (millis() - start < 1500) {
    if (unoSerial.available()) {
      String reply = unoSerial.readStringUntil('\n');
      reply.trim();
      return reply;
    }
  }
  return "ERROR|TIMEOUT";
}

void handleCheck() {
  if (!server.hasArg("item") || !server.hasArg("size")) {
    server.send(400, "text/plain", "ERROR|MISSING_PARAMETER");
    return;
  }
  String command = "GET|" + server.arg("item") + "|" + server.arg("size");
  server.send(200, "text/plain", askUno(command));
}

void handleUpdate() {
  if (!server.hasArg("item") || !server.hasArg("size") || !server.hasArg("qty")) {
    server.send(400, "text/plain", "ERROR|MISSING_PARAMETER");
    return;
  }
  String command = "SET|" + server.arg("item") + "|" +
                   server.arg("size") + "|" + server.arg("qty");
  server.send(200, "text/plain", askUno(command));
}

void setup() {
  WiFi.softAP("Uniform-Kiosk", "school123");
  unoSerial.begin(9600, SERIAL_8N1, 16, 17); // RX2, TX2

  server.on("/", []() { server.send(200, "text/html", kioskHtml); });
  server.on("/check", handleCheck);
  server.on("/update", handleUpdate);
  server.on("/status", handleStatus);
  server.on("/assist", handleAssist);
  server.begin();
}

void loop() {
  server.handleClient();
  readUnsolicitedArduinoEvents();
}'''
add_code(doc, esp_code, "Simplified ESP32 program structure", size=7.7)
add_para(doc, "The request /check?item=PE_SHIRT&size=M becomes GET|PE_SHIRT|M. The /update route should require a valid administrator session or prototype token before it sends a SET command. The sample route is simplified and must not be described as secure.")
add_para(doc, "A basic implementation may store the HTML in a program string on the ESP32. A later version may use the ESP32 file system, but that is optional. Keep the first version small enough to debug.")

# Section 27
add_heading(doc, "Section 27 HTML Structure", 1)
add_para(doc, "HTML describes the content of the kiosk. The project may use one HTML page with several screen sections. JavaScript hides all sections except the active one. This approach is easier than building separate pages for a small prototype.")
html_code = r'''<main id="kiosk">
  <section id="home" class="screen active">...</section>
  <section id="uniform-screen" class="screen">...</section>
  <section id="size-screen" class="screen">...</section>
  <section id="confirm-screen" class="screen">...</section>
  <section id="result-screen" class="screen">...</section>
  <section id="admin-login" class="screen">...</section>
  <section id="admin-panel" class="screen">...</section>
</main>'''
add_code(doc, html_code, "One-page screen structure", size=8.5)
add_para(doc, "Each button should have a meaningful label and a predictable click action. Avoid links or controls that lead to unfinished screens. Add a short loading message while waiting for the Arduino response so the student does not press the button repeatedly.")

# Section 28
add_heading(doc, "Section 28 CSS Design", 1)
add_para(doc, "CSS controls the appearance of the kiosk. Use a high-contrast design with readable text, generous spacing, and consistent button shapes. A button height of about 60 to 80 pixels is a useful starting point for a touchscreen, but the final size should be tested on the actual display.")
css_code = r'''body {
  margin: 0;
  font-family: Arial, sans-serif;
  background: #f4f6f8;
  color: #1f2933;
}

.screen {
  display: none;
  max-width: 900px;
  margin: 0 auto;
  padding: 32px;
}

.screen.active { display: block; }

button {
  min-height: 68px;
  min-width: 180px;
  margin: 10px;
  padding: 12px 20px;
  font-size: 22px;
  border-radius: 8px;
}

button:focus { outline: 4px solid #17365d; }'''
add_code(doc, css_code, "Simple touchscreen-friendly CSS", size=8.2)
add_para(doc, "Use status text as well as color. Test the interface from the normal viewing distance. Do not claim formal accessibility compliance unless the interface has been evaluated against a recognized standard.")

# Section 29
add_heading(doc, "Section 29 JavaScript Function", 1)
add_para(doc, "JavaScript remembers the selected uniform and size, changes the visible screen, sends requests to the ESP32, reads the reply, and updates the result. It should prevent a check until both selections are present.")
js_code = r'''let selectedItem = "";
let selectedSize = "";

async function checkAvailability() {
  if (!selectedItem || !selectedSize) return;
  showScreen("loading-screen");
  try {
    const url = `/check?item=${encodeURIComponent(selectedItem)}` +
                `&size=${encodeURIComponent(selectedSize)}`;
    const response = await fetch(url);
    const reply = await response.text();
    if (!response.ok || reply.startsWith("ERROR|")) throw new Error(reply);
    const [type, quantity, status] = reply.split("|");
    if (type !== "STOCK") throw new Error("Unexpected response");
    displayResult(Number(quantity), status);
    showScreen("result-screen");
  } catch (error) {
    displayError("The hardware controller did not return a valid result.");
    showScreen("error-screen");
  }
}'''
add_code(doc, js_code, "Example availability request", size=8.2)
add_para(doc, "The script should also reset old selections when Home is pressed, display current choices on the confirmation screen, validate administrator quantities, and poll or request assistance status if the Arduino sends an assistance event.")

# Section 30
add_heading(doc, "Section 30 Sample Database Structure", 1)
add_para(doc, "The prototype does not need a database server. Its conceptual database is a 6 by 6 stock table stored in Arduino memory and EEPROM. Each item and size pair maps to one address.")
add_code(doc, """PE Shirt
XS = 10   S = 15   M = 12   L = 8   XL = 4   XXL = 2

Male Polo
XS = 12   S = 8    M = 5    L = 3   XL = 2   XXL = 0

EEPROM address = (item index * 6) + size index""", "Conceptual record structure")
add_table(doc, ["Item index", "Uniform", "EEPROM addresses"], [
    (0, "Male Polo", "0 to 5"), (1, "Female Blouse", "6 to 11"),
    (2, "Male Pants", "12 to 17"), (3, "Female Skirt", "18 to 23"),
    (4, "PE Shirt", "24 to 29"), (5, "PE Pants", "30 to 35"),
], widths=[1.2, 3.0, 2.5], aligns=["center", "left", "center"])
add_para(doc, "Within each six-address block, XS is offset 0, S is 1, M is 2, L is 3, XL is 4, and XXL is 5. Keep this order unchanged in the code and documentation.")
# Section 31
add_heading(doc, "Section 31 Functional Testing", 1)
add_para(doc, "Run each test on the assembled prototype. Write the observed result in Actual Result and mark Pass only when the display, serial response, and hardware behavior match the expected result. A blank row means the test has not yet been performed.")
test_rows = [
    ("T01", "Available stock check", "PE Shirt M with stock 12", "Display 12 and Available; green LED on", "", ""),
    ("T02", "Low stock check", "Male Polo XL with stock 2", "Display 2 and Low Stock; yellow LED on", "", ""),
    ("T03", "Out of stock check", "Male Polo XXL with stock 0", "Display 0 and Out of Stock; red LED on", "", ""),
    ("T04", "No automatic deduction", "Check PE Shirt M twice", "Both checks display 12", "", ""),
    ("T05", "Assistance button", "Press physical button once", "Short buzzer sound and assistance message", "", ""),
    ("T06", "Admin update", "Change PE Shirt M from 12 to 15", "Confirmation received; next check displays 15", "", ""),
    ("T07", "Power recovery", "Restart Uno after saved update", "Updated value remains available", "", ""),
    ("T08", "Invalid quantity", "Enter -1 or blank", "Update rejected; old value unchanged", "", ""),
    ("T09", "Invalid item code", "Send GET|UNKNOWN|M", "Error response; no false stock result", "", ""),
    ("T10", "Arduino disconnected", "Check item with Uno link removed", "Timeout or controller error shown", "", ""),
    ("T11", "LED exclusivity", "Run all three stock states", "Only the matching LED is on after each check", "", ""),
    ("T12", "Navigation", "Check item then press Home", "Home opens and old selection is cleared", "", ""),
]
add_table(doc, ["Test", "Feature", "Input", "Expected Result", "Actual Result", "Status"], test_rows, widths=[.55, 1.1, 1.35, 2.15, 1.0, .55], font_size=7.7, aligns=["center", "left", "left", "left", "left", "center"])
add_para(doc, "Also record the date, tester, program versions, and wiring version in the group test log. If a test fails, note the cause and repeat the test after the correction.")

# Section 32
add_heading(doc, "Section 32 Possible Problems and Troubleshooting", 1)
trouble = [
    ("ESP32 will not connect", "Wrong network name or password; weak signal; wrong mode", "Recheck credentials or use ESP32 access-point mode; view Serial Monitor messages."),
    ("Arduino receives no commands", "TX and RX not crossed; wrong baud rate; loose wire", "Connect TX to RX, match baud rates, inspect newline handling, and test with Serial Monitor."),
    ("Wrong stock value", "Item index or EEPROM address mismatch", "Print item and size indices; verify the 6 by 6 mapping; reinitialize test data safely."),
    ("LED always on", "Pin logic error or LED connected incorrectly", "Turn all LEDs off before setting one; check polarity and pin numbers."),
    ("Button always pressed", "INPUT_PULLUP wiring wrong or pin shorted to GND", "Confirm one side goes to D2 and the other to GND; inspect the breadboard rows."),
    ("ESP32 resets", "Unstable power supply, short circuit, or excessive current", "Use a stable rated supply; inspect wiring; do not power heavy loads from GPIO."),
    ("Webpage cannot connect", "Computer is on another network or wrong ESP32 IP", "Connect to the kiosk network and open the IP printed by the ESP32."),
    ("Serial text is garbled", "Baud rates differ or shared ground is missing", "Use the same baud rate and connect common GND."),
    ("TX and RX wiring fails", "TX connected to TX or RX connected to RX", "Cross the signals: ESP32 TX to Uno RX and Uno TX to ESP32 RX through safe level conversion."),
    ("Missing common ground", "Boards have different voltage references", "Connect Arduino GND and ESP32 GND."),
    ("5 V reaches ESP32 GPIO", "Uno TX connected directly", "Disconnect power immediately and add a level converter or verified divider."),
    ("EEPROM value is 255 or corrupted", "Memory not initialized or invalid data stored", "Use a first-run marker and range checks; restore known sample values."),
    ("Update says success but value is old", "ESP32 confirms before Arduino reply", "Show success only after receiving UPDATED from the Arduino."),
    ("One button creates many alerts", "Mechanical button bounce", "Add debounce logic and trigger only on the HIGH-to-LOW transition."),
]
add_table(doc, ["Problem", "Possible cause", "Simple solution"], trouble, widths=[1.55, 2.25, 2.9], font_size=8.3, aligns=["left", "left", "left"])
add_warning(doc, "Safety first.", "Remove power before moving wires. If a board becomes hot, smells unusual, or repeatedly resets, disconnect it and ask a teacher or electronics adviser to inspect the circuit.")

# Section 33
add_heading(doc, "Section 33 Development Order", 1)
phases = [
    (1, "Arduino LEDs and button", "Blink each LED separately, read the button, and sound the buzzer safely."),
    (2, "Arduino stock arrays", "Store the 36 sample values and print selected quantities to Serial Monitor."),
    (3, "Arduino GET command", "Parse a GET command and return the correct quantity and status."),
    (4, "Arduino SET command", "Validate and change one value in RAM before adding EEPROM."),
    (5, "ESP32 serial connection", "Send fixed commands between boards and confirm clean replies."),
    (6, "ESP32 web server", "Open a basic page and test /check from a browser."),
    (7, "HTML interface", "Build all student screens with working navigation using sample results."),
    (8, "JavaScript integration", "Connect the interface to /check and handle errors and loading states."),
    (9, "Admin interface", "Add PIN screen, current stock lookup, validation, and /update."),
    (10, "EEPROM storage", "Save changed values and verify them after a power restart."),
    (11, "Physical kiosk enclosure", "Mount only after the electronics and software work on the bench."),
    (12, "Testing", "Complete the test table, correct failures, and rehearse the demonstration."),
]
add_table(doc, ["Phase", "Build step", "Completion check"], phases, widths=[.7, 1.7, 4.3], font_size=8.8, aligns=["center", "left", "left"])
add_warning(doc, "Do not skip phases.", "Make each phase work by itself before adding the next. When several new parts are added at once, it becomes difficult to identify the cause of a failure.")

# Section 34
add_heading(doc, "Section 34 Demonstration Scenario", 1)
add_numbered(doc, [
    "A student selects PE Shirt and Medium.",
    "The screen displays 12 pieces available.",
    "The green LED turns on.",
    "An administrator opens the admin screen and changes the value to 2.",
    "The Arduino confirms the update and stores it.",
    "The student checks the same item again.",
    "The screen displays Low Stock with 2 pieces.",
    "The yellow LED turns on and the green LED turns off.",
    "The administrator changes the value to 0.",
    "The student checks the same item again.",
    "The screen displays Out of Stock.",
    "The red LED turns on and the other LEDs turn off.",
    "The student presses the physical assistance button.",
    "The buzzer sounds and the assistance message appears.",
])
add_para(doc, "This sequence proves that the browser sends a request, the ESP32 forwards it, the Arduino applies the rule, the hardware output changes, the admin update changes the saved quantity, and the physical button can create an event.")
# Section 35
add_heading(doc, "Section 35 Questions the Panel May Ask", 1)
qa = [
    ("Why did you use Arduino Uno?", "The Uno stores or accesses prototype stock values, processes GET and SET commands, applies the stock rule, controls the LEDs and buzzer, and reads the assistance button."),
    ("Why did you also use ESP32?", "The ESP32 provides Wi-Fi and a web server. It connects the browser interface to the Arduino through serial communication."),
    ("Why did you not use ESP32 only?", "The project is designed to demonstrate an Arduino-based hardware controller. Separating the Uno hardware logic from the ESP32 communication role also makes the architecture easy to explain."),
    ("Why is stock manually updated?", "This is a limited prototype without a link to the school's official inventory or sales process. Authorized staff enter a verified quantity."),
    ("Does checking stock decrease inventory?", "No. A check is read-only. Stock changes only after an authorized SET update."),
    ("Why is there no purchasing feature?", "The project's problem is availability checking. Purchasing would add transaction, receipt, security, and policy requirements outside the prototype scope."),
    ("Why is there no payment system?", "Payment handling would require financial security, account controls, audit processes, and approval. It is unnecessary for proving the stock-checking concept."),
    ("Where is the inventory stored?", "The working values are in the Arduino stock array, and saved quantities are kept in EEPROM for restart recovery."),
    ("What happens if the kiosk loses power?", "EEPROM keeps the last saved quantities. The Arduino loads them when it starts again. Unsaved changes would not be preserved."),
    ("How do you determine Low Stock?", "A quantity from 1 to 3 is Low Stock. Four or more is Available, and zero is Out of Stock."),
    ("What is the purpose of the LEDs?", "They provide a visible physical output that matches the status calculated by the Arduino."),
    ("What makes the kiosk interactive?", "The user chooses an item and size, the system processes the request, hardware responds, and the screen changes based on the returned result."),
    ("Is it ready for actual school deployment?", "No. It is a prototype. Deployment would require verified inventory integration, stronger security, reliability testing, maintenance procedures, and school approval."),
]
add_table(doc, ["Question", "Suggested answer"], qa, widths=[2.0, 4.7], font_size=8.8, aligns=["left", "left"])

# Section 36
add_heading(doc, "Section 36 Future Improvements", 1)
add_para(doc, "Possible future developments include a central database, cloud synchronization, barcode scanning, RFID inventory tracking, connection to the school's official inventory, automatic stock deduction after an approved transaction, support for multiple kiosks, staff notifications, purchasing integration, and a mobile application.")
add_warning(doc, "Future developments only.", "These functions are not part of the current prototype and should not be shown as completed features. Add them only after the required stock-checking system has been tested and after the school approves the added data and security responsibilities.")

# Section 37
add_heading(doc, "Section 37 Final Recommended System", 1)
add_diagram(doc, [
    "TOUCHSCREEN OR COMPUTER", "|", "v", "HTML + CSS + JAVASCRIPT", "|", "v",
    "ESP32 WEB SERVER", "|", "v", "SERIAL COMMUNICATION", "|", "v", "ARDUINO UNO",
    "|", "v", "STOCK LOGIC + EEPROM", "|", "v",
    "GREEN LED + YELLOW LED + RED LED", "+", "BUZZER", "+", "ASSISTANCE BUTTON",
], "Recommended prototype architecture")
add_heading(doc, "Student Capabilities", 2)
add_bullets(doc, ["Browse uniforms.", "Select a size.", "Check stock quantity and status.", "See the matching LED.", "Request local assistance."])
add_heading(doc, "Administrator Capabilities", 2)
add_bullets(doc, ["Log in using a prototype PIN.", "View the current stored quantity.", "Update and save a stock quantity."])
add_heading(doc, "Functions Not Included", 2)
add_bullets(doc, ["Buying or online ordering.", "Payments or financial accounts.", "Reservations.", "Automatic actual inventory transactions.", "Automatic dispensing or enterprise inventory management."])
add_para(doc, "The recommended design keeps the capstone focused and demonstrable. The computer provides the interface, the ESP32 carries web requests, and the Arduino Uno remains responsible for the stock rule and physical hardware. A successful prototype should return correct quantities, activate the correct LED, save authorized updates, detect the assistance button, and show clear errors when communication fails.")

# Final checklist
add_heading(doc, "Prototype Completion Checklist", 1)
checklist = [
    "All six uniform types and six sizes can be selected.",
    "All 36 stock values return the expected quantity.",
    "Available, Low Stock, and Out of Stock use the same thresholds everywhere.",
    "Only the correct status LED turns on.",
    "Checking an item never reduces stock.",
    "Admin updates are validated, confirmed by Arduino, and retained after restart.",
    "The assistance button creates one buzzer event and one message per press.",
    "The Uno-to-ESP32 voltage level is protected and both devices share ground.",
    "Timeouts and invalid commands show errors instead of false results.",
    "The completed test table contains actual results and dates.",
    "The group can explain each device's role and demonstrate the full data path.",
]
for item in checklist:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    r = p.add_run("[ ]  " + item)
    set_font(r)

# Core properties and field update hint
doc.core_properties.title = TITLE
doc.core_properties.subject = "Grade 12 capstone system development guide"
doc.core_properties.author = "Grade 12 Capstone Project Team"
settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

doc.save(OUTPUT)
print(OUTPUT)
