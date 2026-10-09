from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "handoff" / "UNIVUE_COMPLETE_AI_HANDOFF.md"
OUTPUT = ROOT / "handoff" / "UNIVUE_COMPLETE_AI_HANDOFF.docx"
LOGO = ROOT / "firmware" / "esp32" / "cutte_kiosk_bridge" / "data" / "images" / "branding" / "nu-shield.png"

BLACK = RGBColor(0, 0, 0)
NAVY = "172B59"
PALE = "F4F6FA"
GRID = "D9D9D9"


def set_font(run, name="Arial", size=None, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, value=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "5")
        tag.set(qn("w:color"), GRID)
        borders.append(tag)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:cantSplit")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_rich_text(paragraph, text, size=10.5):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", max(8.5, size - 1), color=BLACK)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size)


def add_table(doc, rows):
    headers = [cell.strip() for cell in rows[0].strip("|").split("|")]
    body = rows[2:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_borders(table)
    repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, NAVY)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(p, header, 8.5)
        for run in p.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for row_index, row in enumerate(body):
        values = [cell.strip() for cell in row.strip("|").split("|")]
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(values):
            cell = cells[index]
            if row_index % 2:
                shade(cell, PALE)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_rich_text(p, value, 8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.keep_together = True

    list_bullet = doc.styles["List Bullet"]
    list_bullet.paragraph_format.keep_together = True

    settings = {
        "Title": (28, 0, 13),
        "Subtitle": (12, 0, 8),
        "Heading 1": (17, 16, 7),
        "Heading 2": (13, 12, 5),
        "Heading 3": (11.5, 9, 4),
    }
    for name, (size, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        p_pr = style._element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.footer_distance = Inches(0.28)
    configure_styles(doc)

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(0.82))
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich_text(title, "UNIVUE Complete AI Project Handoff", 28)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich_text(subtitle, "NU Baliwag Bulldogs Exchange", 12)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich_text(date, "Current architecture and transfer guide  |  10 October 2026", 10)
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_before = Pt(16)
    intro.paragraph_format.left_indent = Inches(0.55)
    intro.paragraph_format.right_indent = Inches(0.55)
    add_rich_text(intro, "Use this manual with the repository when transferring UNIVUE to another AI account, developer, or group member. It identifies the current source of truth and separates the live Supabase and Raspberry Pi system from the older Arduino and ESP32 prototype.", 11)
    doc.add_page_break()

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    index = 1
    in_code = False
    code_lines = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("# "):
            index += 1
            continue
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.space_after = Pt(5)
                for code_index, code in enumerate(code_lines):
                    if code_index:
                        p.add_run().add_break()
                    run = p.add_run(code or " ")
                    set_font(run, "Consolas", 8.5)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            index += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            index += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            index += 1
            continue
        if line.startswith("| ") and index + 1 < len(lines) and re.match(r"^\|[\s:|\-]+\|$", lines[index + 1]):
            table_rows = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                table_rows.append(lines[index])
                index += 1
            add_table(doc, table_rows)
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(5)
            while index < len(lines) and re.match(r"^\d+\. ", lines[index]):
                match = re.match(r"^(\d+)\. (.*)$", lines[index])
                if len(p.runs):
                    p.add_run().add_break()
                add_rich_text(p, f"{match.group(1)}. {match.group(2)}")
                index += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, line[2:])
            index += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.right_indent = Inches(0.28)
            add_rich_text(p, line[2:], 10.5)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.startswith("The `Need help` button"):
            doc.add_page_break()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.first_line_indent = Inches(0)
        add_rich_text(p, line)
        index += 1

    doc.core_properties.title = "UNIVUE Complete AI Project Handoff"
    doc.core_properties.subject = "Current architecture operations security hardware and transfer guide"
    doc.core_properties.author = "UNIVUE Project Team"
    doc.core_properties.keywords = "UNIVUE, NU Baliwag, Supabase, Raspberry Pi, PayMongo, handoff"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
